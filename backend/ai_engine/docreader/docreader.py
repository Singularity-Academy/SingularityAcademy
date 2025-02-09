import fitz  # PyMuPDF
import docx
import csv
from pptx import Presentation

class DocReader:
    def __init__(self, file_path):
        self.file_path = file_path

    def read_pdf(self):
        doc = fitz.open(self.file_path)
        text = ""
        for page_num in range(doc.page_count):
            page = doc.load_page(page_num)
            text += page.get_text()
        return text

    def read_docx(self):
        doc = docx.Document(self.file_path)
        text = ""
        for paragraph in doc.paragraphs:
            text += paragraph.text + '\n'
        return text

    def read_ppt(self):
        presentation = Presentation(self.file_path)
        text = ""
        for slide in presentation.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text += shape.text + '\n'
        return text

    def read_csv(self):
        text = ""
        with open(self.file_path, mode='r', encoding='utf-8') as file:
            reader = csv.reader(file)
            for row in reader:
                text += ",".join(row) + '\n'
        return text

    def read(self):
        file_extension = self.file_path.split('.')[-1].lower()
        if file_extension == 'pdf':
            return self.read_pdf()
        elif file_extension == 'docx':
            return self.read_docx()
        elif file_extension == 'pptx':
            return self.read_ppt()
        elif file_extension == 'csv':
            return self.read_csv()
        else:
            return "Unsupported file format."